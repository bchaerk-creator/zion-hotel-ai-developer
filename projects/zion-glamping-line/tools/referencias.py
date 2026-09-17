# -*- coding: utf-8 -*-
"""Referências de mercado da linha: cotações NOMASTRA (Guangzhou, 1/mai/2025, contato Nancy) para os lodges H23 (5,2 x 5,2 m,
referência da ZION LODGE 24) e H28 (3,8 x 7,8 m, referência da ZION LODGE 28). O comparativo técnico (sem preços) entra no
catálogo, na apresentação e na especificação; os preços (USD EXW) ficam apenas nos arquivos internos:
interno/ZION_REFERENCIAS_NOMASTRA_H23_H28.html/.pdf/.xlsx.  Uso: python3 referencias.py"""
import os, html
ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
INT = os.path.join(ROOT, "interno")
FORNECEDOR = "NOMASTRA Co., Ltd. · No.7, East 2nd Road, Hengli Town, Nansha District, Guangzhou, China"
CONTATO = "Nancy · nancyli@nomastra.com · +86 181 2794 8011"
COMERCIAL = "EXW · TT 50 % no pedido + 50 % antes do embarque · 15 a 20 dias úteis após o sinal · garantia de 1 ano · engenheiro de montagem a US$ 150/dia + visto, passagens, hotel e transporte"

# premissas internas para custo posto (ajustar a cada cotação)
FX = 5.50                 # R$/US$ de referência
FRETE_USD = 6500.0        # frete marítimo + seguro + THC, 1 contêiner 40' HC (a cotar); os dois kits juntos cabem em um contêiner
TRIBUTOS = 0.62           # II + IPI + PIS/COFINS + ICMS sobre o CIF (a confirmar com despachante; NCM 9406 estruturas pré-fabricadas)
DESEMBARACO_BRL = 9000.0  # despachante, armazenagem, transporte interno até SC

# ---------------------------------------------------------------------------- comparativos comuns (o que é igual nos dois kits)
_ESTRUTURA_REF = "tubo de aço 80 x 80 com pintura dupla e revestimento de alumínio · base galvanizada a fogo · fixações inox"
_COBERTURA_REF = "membrana PVDF 1050 g/m² (B1) · segunda camada PVC 850 g/m² bege"
_PAREDES_REF = "compósito de cimento espumado 80 mm com lâmina de madeira (100 mm no total)"
_BANHO_REF = "módulo pré-fabricado 2,6 x 1,5 x 2,2 m (honeycomb, ACM, bandeja ABS), opcional"
_INST_REF = "tubos e fiação pré-instalados (opcional) · sem climatização"
_FUND_REF = "deck em quadro de aço sobre o terreno (20 a 35 cm)"
_ORIGEM_REF = "Guangzhou, China · EXW · 15 a 20 dias úteis + trânsito marítimo e desembaraço · engenheiro de montagem cobrado à parte"

REFS = {
    "lodge24": dict(
        modelo="H23", zion="ZION LODGE 24", pdf="interno/referencias/NOMASTRA_H23_Lodge_Tent_Quotation_2025-05.pdf", data="1 de maio de 2025",
        tamanho="5,20 x 5,20 m", area="23 m²", terraco="sem terraço", altura="5,4 m", beiral="2,6 m", portas=1, janelas=2,
        itens=[("Body Tent (estrutura, base, membranas, forro elástico, paredes, porta e janelas)", 12280.00, "padrão"),
               ("Deck de fundação 45 m² (quadro de aço galvanizado + madeira tratada, 20 a 35 cm)", 2700.00, "opcional"),
               ("Piso interno SPC 6 mm", 545.00, "opcional"),
               ("Módulo de banho 2,6 x 1,5 x 2,2 m (quadro galvanizado, honeycomb, ACM, box 5 mm, louças, ducha, exaustor, LED, bandeja ABS)", 3000.00, "opcional"),
               ("Tubulações e fiação pré-instaladas (layout antes do embarque)", 840.00, "opcional"),
               ("Cortinas blackout com trilho", 1230.00, "opcional"),
               ("Mobiliário e luminárias (cama 1,80 x 2,00, 4 criados, 2 conjuntos sofá + mesa, guarda-roupa, fita LED, 2 arandelas, 1 pendente)", 3080.00, "opcional")],
        comparativo=[
            ("Planta e área", "5,20 x 5,20 m · 23 m² internos", "octógono regular de 5,40 m entre faces · 24,2 m² internos"),
            ("Terraço / deck", "sem terraço; deck de fundação opcional de 45 m² (20 a 35 cm)", "deck de 6,1 m² em uma face com vela de sombra · deck elevado sobre estacas"),
            ("Alturas", "cume 5,4 m · beiral 2,6 m · cume cego", "Lanterna Zion Ø1,20 de 4,20 a 4,70 m · beiral 2,6 m · luz zenital sobre a cama"),
            ("Estrutura", _ESTRUTURA_REF, "8 pilares Ø101,6 revestidos em madeira · anel de beiral 150 x 100 · 8 caibros Ø76 · anel de compressão · tudo galvanizado a fogo, kit parafusado"),
            ("Cobertura", _COBERTURA_REF, "membrana PVDF 1050 g/m² em gomos com keder · câmara ventilada 60 mm · lã de PET 50 mm · forro tensionado acústico"),
            ("Forro", "tecido elástico 140 g/m², sem isolamento", "forro tensionado sob a membrana + painel ripado termotratado nas faces opacas"),
            ("Paredes", _PAREDES_REF, "painel SIP 90 mm + ripado termotratado em cinco faces (privacidade em implantações densas)"),
            ("Portas", "1 porta simples de vidro duplo · 5 + 20 A + 5 mm", "PC1 porta de correr 2,00 x 2,40 · vidro insulado 6 lam + 12 Ar + 6 temp low-e · alumínio bronze RPT"),
            ("Janelas", "2 janelas de vidro duplo 5 + 20 A + 5 mm", "três faces inteiras de vidro insulado voltadas à paisagem + janela do café + fresta do banho + lanterna"),
            ("Piso", "SPC 6 mm (opcional)", "carvalho de engenharia 14 mm · porcelanato no banho · lã de PET sob o piso"),
            ("Banho", _BANHO_REF, "banho integrado de 5,0 m² no segmento posterior: bancada em pedra, chuveiro"),
            ("Instalações", _INST_REF, "elétrica, hidráulica, climatização split, fechadura digital e cenas"),
            ("Fundação", _FUND_REF, "15 estacas helicoidais + deck e piso em vigas U 150; sem concreto"),
            ("Mobiliário", "cama 1,80 x 2,00, 4 criados, 2 conjuntos sofá + mesa, guarda-roupa, fita LED, 2 arandelas, 1 pendente (opcional)", "FF&E Zion New Luxury item a item (cama king, criados, duas poltronas, closet, café/minibar, luminárias, enxoval) · sem pendentes"),
            ("Origem e prazo", _ORIGEM_REF, "fábrica em Santa Catarina · kit parafusado · montagem em 6 dias pela equipe Zion · sem importação")]),
    "lodge28": dict(
        modelo="H28", zion="ZION LODGE 28", pdf="interno/referencias/NOMASTRA_H28_Lodge_Tent_Quotation_2025-05.pdf", data="1 de maio de 2025",
        tamanho="3,80 x 7,80 m", area="26 m²", terraco="terraço 18 m²", altura="5,0 m", beiral="2,6 m", portas=2, janelas=2,
        itens=[("Body Tent (estrutura, base, membranas, forro rígido, paredes, portas e janelas)", 17410.00, "padrão"),
               ("Deck de fundação 50 m² (quadro de aço galvanizado + madeira tratada, 20 a 35 cm)", 3000.00, "opcional"),
               ("Piso interno SPC 6 mm", 615.00, "opcional"),
               ("Módulo de banho 2,6 x 1,5 x 2,2 m (quadro galvanizado, honeycomb, ACM, box 5 mm, louças, ducha, exaustor, LED, bandeja ABS)", 3000.00, "opcional"),
               ("Tubulações e fiação pré-instaladas (layout antes do embarque)", 948.00, "opcional"),
               ("Cortinas blackout com trilho", 1400.00, "opcional"),
               ("Mobiliário e luminárias (cama 1,80 x 2,00, 2 criados, sofá + mesa, 2 cadeiras externas, guarda-roupa, fita LED, 2 arandelas, 2 pendentes)", 3895.00, "opcional")],
        comparativo=[
            ("Planta e área", "3,80 x 7,80 m · 26 m² internos", "octógono alongado 4,20 x 7,80 m · 29,7 m² internos"),
            ("Terraço / deck", "terraço 18 m²; deck de fundação opcional de 50 m² (20 a 35 cm)", "terraço 17,5 m² em três faces com vela de sombra · deck elevado sobre estacas"),
            ("Alturas", "cume 5,0 m · beiral 2,6 m · cume cego", "duas Lanternas Zion Ø1,10 de 4,00 a 4,45 m · beiral 2,6 m · luz zenital sobre cama e estar"),
            ("Estrutura", _ESTRUTURA_REF, "8 pilares Ø101,6 revestidos em madeira · anel de beiral 150 x 100 · 10 caibros Ø76 · anéis de compressão · tudo galvanizado a fogo, kit parafusado"),
            ("Cobertura", _COBERTURA_REF, "membrana PVDF 1050 g/m² em gomos com keder · câmara ventilada 60 mm · lã de PET 50 mm · forro tensionado acústico"),
            ("Forro", "estrutura de aço + painel decorativo de fibra de bambu e madeira", "forro tensionado sob a membrana + painel ripado termotratado nas faces opacas"),
            ("Paredes", _PAREDES_REF, "painel SIP 90 mm + ripado termotratado em três faces (banho e cabeceira)"),
            ("Portas", "1 porta de correr em vidro duplo + 1 porta simples de vidro · 5 + 20 A + 5 mm", "PC1 porta de correr 2,00 x 2,40 · vidro insulado 6 lam + 12 Ar + 6 temp low-e · alumínio bronze RPT"),
            ("Janelas", "2 janelas de vidro duplo 5 + 20 A + 5 mm", "cinco faces inteiras de vidro insulado + fresta alta do banho + duas lanternas"),
            ("Piso", "SPC 6 mm (opcional)", "carvalho de engenharia 14 mm · porcelanato no banho · lã de PET sob o piso"),
            ("Banho", _BANHO_REF, "banho integrado de 6,2 m² no segmento posterior: bancada em pedra, chuveiro, fresta alta"),
            ("Instalações", _INST_REF, "elétrica, hidráulica, climatização dutada (evaporadora no ático), automação de cenas"),
            ("Fundação", _FUND_REF, "21 estacas helicoidais + deck e piso em vigas U 150; sem concreto"),
            ("Mobiliário", "cama 1,80 x 2,00 e criados em nogueira, sofá em acácia, 2 cadeiras externas, guarda-roupa, fita LED, 2 arandelas, 2 pendentes (opcional)", "FF&E Zion New Luxury item a item (cama king, criados, sofá, closet, café/minibar, luminárias, enxoval, deck) · sem pendentes"),
            ("Origem e prazo", _ORIGEM_REF, "fábrica em Santa Catarina · kit parafusado · montagem em 8 dias pela equipe Zion · sem importação")]),
}

def comparativo(code): return REFS[code]["comparativo"]
def ref(code): return REFS[code]
def money_usd(v): return "US$ " + f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
def money_brl(v): return "R$ " + f"{v:,.0f}".replace(",", ".")

def custo_posto(code, frete_usd=None):
    exw = sum(v for _, v, _ in REFS[code]["itens"]); frete = FRETE_USD if frete_usd is None else frete_usd; cif = exw + frete
    trib = cif * TRIBUTOS; posto = cif * FX + trib * FX + DESEMBARACO_BRL
    return dict(exw_usd=exw, frete_usd=frete, cif_usd=cif, cif_brl=cif * FX, tributos_brl=trib * FX, desembaraco_brl=DESEMBARACO_BRL, posto_brl=posto)

def build():
    os.makedirs(INT, exist_ok=True); esc = html.escape
    try:
        from build_catalogo import cost
        est = {c: cost(c) for c in REFS}
    except Exception:
        est = {}
    secs = ""
    for code, R in REFS.items():
        cp = custo_posto(code)
        itens = "".join(f"<tr><td>{esc(d)}</td><td>{esc(k)}</td><td class='num'>{money_usd(v)}</td></tr>" for d, v, k in R["itens"])
        comp = "".join(f"<tr><th>{esc(a)}</th><td>{esc(b)}</td><td>{esc(c)}</td></tr>" for a, b, c in R["comparativo"])
        e = est.get(code); zion = f"<p>Estimativa paramétrica interna da {R['zion']} (1ª unidade, Zion Standard, sem FF&amp;E): <b>{money_brl(e[0])}</b>; 10ª unidade ≈ {money_brl(e[2])}.</p>" if e and e[0] else ""
        secs += f"""<h1>REFERÊNCIA DE MERCADO · NOMASTRA {esc(R['modelo'])} × {esc(R['zion'])}</h1><small>Cotação recebida em {esc(R['data'])} · fonte: {esc(R['pdf'])}</small>
<h2>COTAÇÃO</h2>
<dl class="kv"><dt>Fornecedor</dt><dd>{esc(FORNECEDOR)}</dd><dt>Contato</dt><dd>{esc(CONTATO)}</dd><dt>Modelo</dt><dd>{esc(R['modelo'])} · {esc(R['tamanho'])} · {esc(R['area'])} internos · {esc(R['terraco'])} · altura {esc(R['altura'])} · beiral {esc(R['beiral'])} · {R['portas']} porta(s) · {R['janelas']} janelas</dd><dt>Condições</dt><dd>{esc(COMERCIAL)}</dd></dl>
<table><tr><th style="width:auto">Item cotado</th><th style="width:70px">Tipo</th><th class="num" style="width:110px">USD / conjunto (EXW)</th></tr>{itens}<tr class="tot"><td>Conjunto completo com todos os opcionais</td><td></td><td class="num">{money_usd(cp['exw_usd'])}</td></tr></table>
<h2>CUSTO POSTO EM SANTA CATARINA (PREMISSAS INTERNAS, AJUSTAR A CADA COTAÇÃO)</h2>
<table class="kvt"><tr><th>EXW completo</th><td class="num">{money_usd(cp['exw_usd'])}</td></tr><tr><th>Frete marítimo + seguro + THC (1 x 40' HC, estimado; um contêiner leva os dois kits, então o rateio reduz este valor)</th><td class="num">{money_usd(cp['frete_usd'])}</td></tr><tr><th>CIF</th><td class="num">{money_usd(cp['cif_usd'])} = {money_brl(cp['cif_brl'])} a R$ {FX:.2f}</td></tr>
<tr><th>Tributos de importação ({TRIBUTOS:.0%} sobre o CIF, II + IPI + PIS/COFINS + ICMS, a confirmar com despachante · NCM 9406)</th><td class="num">{money_brl(cp['tributos_brl'])}</td></tr><tr><th>Desembaraço, armazenagem e transporte interno até SC</th><td class="num">{money_brl(DESEMBARACO_BRL)}</td></tr>
<tr class="tot"><td>Custo posto estimado, sem fundação em estacas, sem montagem, sem climatização, sem enxoval</td><td class="num">{money_brl(cp['posto_brl'])}</td></tr></table>
{zion}
<h2>COMPARATIVO TÉCNICO (versão sem preços publicada no catálogo e na apresentação)</h2>
<table><tr><th>Item</th><th style="width:40%">Referência {esc(R['modelo'])} (NOMASTRA)</th><th style="width:44%">{esc(R['zion'])}</th></tr>{comp}</table>
<div class="pb"></div>"""
    doc = f"""<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><title>Referências de mercado · NOMASTRA H23 e H28 · interno</title><style>
@page{{size:A4;margin:14mm}} body{{font-family:'Aventa','DM Sans',Helvetica,Arial,sans-serif;color:#1B2117;background:#FEF5F0;margin:0;padding:16mm;font-size:11px;line-height:1.45}}
h1{{font-size:20px;letter-spacing:.22em;font-weight:300;margin:0 0 4px}} h2{{font-size:12px;letter-spacing:.25em;margin:22px 0 8px;color:#8B714E}} .warn{{background:#1B2117;color:#FEF5F0;padding:8px 12px;font-size:10px;letter-spacing:.08em;margin-bottom:18px}}
table{{width:100%;border-collapse:collapse;margin:6px 0}} th,td{{border-bottom:1px solid #DED6BF;padding:5px 6px;text-align:left;vertical-align:top}} th{{font-weight:600;width:16%}} table.kvt th{{width:72%}} .num{{text-align:right;white-space:nowrap}} tr.tot td{{font-weight:700;border-top:2px solid #1B2117}}
.kv dt{{float:left;clear:left;width:150px;color:#8B714E;font-size:9px;letter-spacing:.15em;text-transform:uppercase;padding-top:2px}} .kv dd{{margin:0 0 4px 160px}} small{{color:#8B714E}} .pb{{page-break-after:always}}
</style></head><body>
<div class="warn">DOCUMENTO INTERNO · CONTÉM PREÇOS · NÃO ENVIAR À FÁBRICA NEM A CLIENTES</div>
<p><small>Leitura geral: as cotações cobrem kits de tenda com paredes de cimento espumado e módulo de banho pré-fabricado, sem isolamento térmico da cobertura, sem climatização, sem fundação em estacas e com montagem por conta do comprador. As unidades Zion são produtos proprietários (lanternas, faces de vidro, vela de sombra, camadas isoladas) fabricados em SC; a comparação de preço deve ser feita no mesmo escopo (posto + fundação + montagem + climatização + FF&amp;E).</small></p>
{secs}</body></html>"""
    out = os.path.join(INT, "ZION_REFERENCIAS_NOMASTRA_H23_H28.html"); open(out, "w", encoding="utf-8").write(doc); print("html ->", out)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook(); wb.remove(wb.active)
    for code, R in REFS.items():
        cp = custo_posto(code); ws = wb.create_sheet(f"Cotação {R['modelo']} (USD)")
        ws.append(["Item cotado", "Tipo", "USD / conjunto (EXW)"])
        for c in (1, 2, 3): ws.cell(1, c).font = Font(bold=True, color="FEF5F0"); ws.cell(1, c).fill = PatternFill("solid", fgColor="1B2117")
        for d, v, k in R["itens"]: ws.append([d, k, v])
        ws.append(["Conjunto completo (EXW)", "", cp["exw_usd"]]); ws.append([]); ws.append(["Premissas de custo posto (ajustar)", "", ""])
        for k, v in (("Câmbio R$/US$", FX), ("Frete + seguro + THC (US$)", cp["frete_usd"]), ("Tributos sobre CIF", TRIBUTOS), ("Desembaraço e transporte interno (R$)", DESEMBARACO_BRL), ("CIF (US$)", cp["cif_usd"]), ("Tributos (R$)", cp["tributos_brl"]), ("Custo posto estimado (R$)", cp["posto_brl"])): ws.append([k, "", v])
        ws.column_dimensions["A"].width = 90; ws.column_dimensions["B"].width = 12; ws.column_dimensions["C"].width = 22
        ws2 = wb.create_sheet(f"Comparativo {R['modelo']}"); ws2.append(["Item", f"Referência {R['modelo']} (NOMASTRA)", R["zion"]])
        for r in R["comparativo"]: ws2.append(list(r))
        ws2.column_dimensions["A"].width = 18; ws2.column_dimensions["B"].width = 70; ws2.column_dimensions["C"].width = 80
    xo = os.path.join(INT, "ZION_REFERENCIAS_NOMASTRA_H23_H28.xlsx"); wb.save(xo); print("xlsx ->", xo)

if __name__ == "__main__":
    build()
