# -*- coding: utf-8 -*-
"""LONA · padrões de corte das lonas externas, dos forros internos e dos revestimentos, por modelo, a partir da geometria
(fonte única: geometry.py). Cada painel é planificado por triangulação sequencial da tira entre duas curvas da superfície
(as costuras ficam sobre os arcos, caibros, anéis ou linhas de emenda), o que dá a geometria "de forma final", sem
compensação: o confeccionista aplica a compensação do tecido (ensaio biaxial) e as sobras de bolsa/keder indicadas.

Saídas por modelo (<modelo>/lona/):
  LN-01_mapa_paineis.svg           mapa dos painéis externos e internos (posição, costuras, códigos)
  LN-1n_<painel>.svg                um padrão por prancha: contorno cotado, malha 0,50 m, bordas por tipo, recortes, tabela de estações
  LN-90_fixacoes.svg                como a lona se prende aos ferros: keder, bolsas, clamps, harpão, cabo de borda
  <TAG>-LON-001_padroes.dxf         todos os padrões em metros (camadas CONTORNO, KEDER, BOLSA, CLAMP, RECORTE, LINHA, TEXTO)
  <TAG>-LON-001_coordenadas.xlsx    tabelas de estações (para plotagem manual ou CNC)
Uso: python3 lona.py [cocoon zenith lodge capsule]"""
import math, os, sys
import numpy as np
from geometry import Cocoon, Zenith, Lodge, Capsule, CocoonSensorial
from svgkit import *
from svgkit import zion_mark, zion_logo

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
KINDS = {  # tipo de borda -> (cor, largura, legenda)
    "keder": ("#1B2117", 3.0, "KEDER Ø10 · desliza no perfil duplo keder do arco / caibro"),
    "keder_base": ("#4E6E8B", 3.0, "KEDER Ø13 · perfil de arremate de base com calha (trilho 100 x 50)"),
    "clamp": ("#8B714E", 3.0, "CLAMP · barra de alumínio 40 x 6 parafusada (M8 @150) sobre EPDM"),
    "bolsa_tubo": ("#B5945A", 4.0, "BOLSA · lona envolve o tubo de borda (Ø60,3) e fecha com fita keder"),
    "cabo": ("#3A3B3A", 3.0, "BOLSA DE CABO · cabo inox Ø12 com esticador, chapa de canto"),
    "harpao": ("#6B7C84", 2.5, "HARPÃO · perfil de alumínio 40 x 20 nos arcos / terças / cabos"),
    "solda": ("#C2734A", 1.6, "SOLDA HF · emenda de fábrica, sobreposição 40 mm"),
    "livre": ("#8B714E", 1.2, "BORDA LIVRE · bainha dupla 30 mm"),
    "junta": ("#3A3B3A", 2.0, "JUNTA · perfil H de alumínio com EPDM (revestimento rígido)"),
    "cristal": ("#4E8BB8", 3.0, "SOLDA HF LONA-CRISTAL · emenda de 40 mm entre a lona opaca e o PVC cristal (linha da terça)"),
    "ziper": ("#C2734A", 3.0, "ZÍPER · zíper YKK #10 duplo na base do cristal (parede que abre) + keder Ø13 no trilho"),
}


# =============================================================================== planificação
def unfold_strip(L, R):
    """planifica a tira entre as polilinhas 3D L e R (mesmo nº de pontos). Devolve (L2, R2) em metros."""
    L = [np.array(p, float) for p in L]; R = [np.array(p, float) for p in R]; n = len(L)
    L2 = [np.array([0.0, 0.0])]; R2 = [np.array([float(np.linalg.norm(R[0] - L[0])), 0.0])]

    def third(A2, B2, dA, dB, side):
        ab = B2 - A2; d = float(np.linalg.norm(ab))
        if d < 1e-9:
            return A2 + np.array([0.0, dA])
        a = (dA ** 2 - dB ** 2 + d ** 2) / (2 * d); h = math.sqrt(max(dA ** 2 - a ** 2, 0.0))
        px = A2 + a * ab / d; perp = np.array([-ab[1], ab[0]]) / d
        c1, c2 = px + h * perp, px - h * perp
        return c1 if np.sign(ab[0] * (c1 - A2)[1] - ab[1] * (c1 - A2)[0]) * side >= 0 else c2
    for j in range(n - 1):
        Ln = third(L2[j], R2[j], float(np.linalg.norm(L[j + 1] - L[j])), float(np.linalg.norm(L[j + 1] - R[j])), +1)
        Rn = third(R2[j], Ln, float(np.linalg.norm(R[j + 1] - R[j])), float(np.linalg.norm(R[j + 1] - L[j + 1])), -1)
        L2.append(Ln); R2.append(Rn)
    # endireita: corda da borda esquerda vertical; origem no canto inferior esquerdo
    v = L2[-1] - L2[0]
    ang = math.atan2(v[1], v[0]) - math.pi / 2 if np.linalg.norm(v) > 1e-9 else 0.0
    c, s = math.cos(-ang), math.sin(-ang)
    rot = lambda p: np.array([c * p[0] - s * p[1], s * p[0] + c * p[1]])
    L2 = [rot(p) for p in L2]; R2 = [rot(p) for p in R2]
    allp = L2 + R2; mx = min(p[0] for p in allp); my = min(p[1] for p in allp)
    return [(float(p[0] - mx), float(p[1] - my)) for p in L2], [(float(p[0] - mx), float(p[1] - my)) for p in R2]


def uv2d(L2, R2, u, v):
    """ponto 2D no padrão: u ao longo da tira (0 → 1, índice fracionário), v de L (0) a R (1)."""
    n = len(L2) - 1; t = min(max(u, 0.0), 1.0) * n; i = min(int(t), n - 1); f = t - i
    lx = L2[i][0] + (L2[i + 1][0] - L2[i][0]) * f; ly = L2[i][1] + (L2[i + 1][1] - L2[i][1]) * f
    rx = R2[i][0] + (R2[i + 1][0] - R2[i][0]) * f; ry = R2[i][1] + (R2[i + 1][1] - R2[i][1]) * f
    return (lx + (rx - lx) * v, ly + (ry - ly) * v)


def path_uv(L2, R2, wps, n=16):
    """caminho fechado no espaço (u, v): lista de vértices (u, v); entre vértices com v constante segue a curva da tira."""
    out = []
    for (u0, v0), (u1, v1) in zip(wps, wps[1:] + wps[:1]):
        k = 40 if abs(u1 - u0) > 1e-9 and abs(v1 - v0) < 1e-9 else n
        for i in range(k):
            t = i / k; out.append(uv2d(L2, R2, u0 + (u1 - u0) * t, v0 + (v1 - v0) * t))
    return out


def poly_area(pts):
    a = 0.0
    for (x1, y1), (x2, y2) in zip(pts, pts[1:] + pts[:1]): a += x1 * y2 - x2 * y1
    return abs(a) / 2


def plen(pts):
    return sum(math.dist(a, b) for a, b in zip(pts[:-1], pts[1:]))


class Pattern:
    def __init__(self, pid, name, model, group, L2, R2, outline=None, notes=None, stations_every=0.5):
        self.id, self.name, self.model, self.group = pid, name, model, group     # group: "externa" | "interna" | "revestimento"
        self.L2, self.R2 = L2, R2
        self.outline = outline or (L2 + R2[::-1])
        self.edges = []      # (kind, pts2d)
        self.holes = []      # (label, pts2d, kind)
        self.lines = []      # (label, pts2d)
        self.notes = notes or []
        self.qty = 1
        self.material = ""
        xs = [p[0] for p in self.outline]; ys = [p[1] for p in self.outline]
        self.w, self.h = max(xs) - min(xs), max(ys) - min(ys)
        self.area = poly_area(self.outline)
        self.lenL, self.lenR = plen(L2), plen(R2)
        self.stations_every = stations_every

    def stations(self):
        """tabela: s ao longo de L (m), xL, yL, xR, yR, largura (distância L-R na mesma estação)."""
        rows = []; acc = [0.0]
        for a, b in zip(self.L2[:-1], self.L2[1:]): acc.append(acc[-1] + math.dist(a, b))
        total = acc[-1]; step = self.stations_every; s = 0.0
        while s <= total + 1e-6:
            i = max(0, min(len(acc) - 2, next((k for k in range(len(acc) - 1) if acc[k + 1] >= s), len(acc) - 2)))
            f = (s - acc[i]) / max(acc[i + 1] - acc[i], 1e-9)
            L = (self.L2[i][0] + (self.L2[i + 1][0] - self.L2[i][0]) * f, self.L2[i][1] + (self.L2[i + 1][1] - self.L2[i][1]) * f)
            R = (self.R2[i][0] + (self.R2[i + 1][0] - self.R2[i][0]) * f, self.R2[i][1] + (self.R2[i + 1][1] - self.R2[i][1]) * f)
            rows.append((s, L[0], L[1], R[0], R[1], math.dist(L, R)))
            s += step
        if total - rows[-1][0] > 0.05:
            rows.append((total, self.L2[-1][0], self.L2[-1][1], self.R2[-1][0], self.R2[-1][1], math.dist(self.L2[-1], self.R2[-1])))
        return rows


def signature(pat, holes=True):
    return (round(pat.w, 2), round(pat.h, 2), round(pat.area, 2), tuple(sorted((h[0], round(poly_area(h[1]), 2)) for h in pat.holes)) if holes else ())


# =============================================================================== CASULO
def cocoon_patterns(C=None):
    C = C or Cocoon(); pats = []
    N = 64
    HT = C.SPINE[2]; SX1, SX2 = C.SPINE[0], C.SPINE[1]

    def arch_curve(x, off=0.0, u0=0.0, u1=1.0, n=N):
        t0, t1 = C.theta_range(x)
        return [C.section_point(x, t0 + (t1 - t0) * (u0 + (u1 - u0) * i / n), off) for i in range(n + 1)]

    def u_of_theta(x, th):
        t0, t1 = C.theta_range(x); return (th - t0) / (t1 - t0)

    def window_uv(w, xa, xb):
        pts = []
        for i in range(48):
            a = 2 * math.pi * i / 48; cs, sn = math.cos(a), math.sin(a)
            r = (abs(cs) ** 1.5 + abs(sn) ** 1.5) ** (-1 / 1.5)
            x = w["xc"] + w["lx"] * r * cs; th = w["tc"] + w["lt"] * r * sn
            pts.append((u_of_theta((xa + xb) / 2, th), (x - xa) / (xb - xa)))
        return pts

    def add_windows(pat, xa, xb, L2, R2, umin=0.0, umax=1.0):
        for w in C.WINDOWS:
            if xa - 0.05 <= w["xc"] <= xb + 0.05:
                uv = window_uv(w, xa, xb)
                if umin <= uv[0][0] <= umax:
                    pat.holes.append((f'{w["name"]} · {2 * w["lx"]:.2f} x {2 * w["lt"] * C.b(w["xc"]):.2f} m'.replace(".", ","), [uv2d(L2, R2, u, v) for (u, v) in uv], "recorte"))

    # ---- painéis externos entre arcos (P1..P7) + cauda (P8)
    bays = [(C.ARCH_X[i], C.ARCH_X[i + 1]) for i in range(7)] + [(C.ARCH_X[7], 9.56)]
    for k, (xa, xb) in enumerate(bays, 1):
        Lc, Rc = arch_curve(xa), arch_curve(xb)
        L2, R2 = unfold_strip(Lc, Rc)
        xm = (xa + xb) / 2; ua = u_of_theta(xm, math.pi / 2 - HT); ub = u_of_theta(xm, math.pi / 2 + HT)
        overlap = max(xa, SX1) < min(xb, SX2)
        full = overlap and SX1 <= xa + 1e-6 and SX2 >= xb - 1e-6
        tail = k == 8
        if full:   # dividido pela Espinha de Luz em duas peças (direita / esquerda)
            clear = getattr(C, "CLEAR", None)
            clear = clear if clear and clear["x1"] - 1e-6 <= xa and xb <= clear["x2"] + 1e-6 else None
            for side, (u0, u1) in (("D", (0.0, ua)), ("E", (ub, 1.0))):
                pieces = [(f"P{k}{side}", u0, u1, "externa", None)]
                if clear:   # cinturão transparente: cristal do trilho de base até z = CLEAR.z; lona acima
                    zc = clear["z"]; t0, t1 = C.theta_range(xm); b = C.b(xm)
                    th = math.asin(min(1.0, max(-1.0, (zc - C.ZC) / b)))
                    uz = (th - t0) / (t1 - t0) if side == "D" else ((math.pi - th) - t0) / (t1 - t0)
                    pieces = [(f"T{k}{side}", u0, uz, "externa", "cristal"), (f"P{k}{side}", uz, u1, "externa", None)] if side == "D" else [(f"P{k}{side}", u0, uz, "externa", None), (f"T{k}{side}", uz, u1, "externa", "cristal")]
                for pid, ua_, ub_, grp, kind in pieces:
                    outline = path_uv(L2, R2, [(ua_, 0), (ub_, 0), (ub_, 1), (ua_, 1)])
                    lab = "CRISTAL (PVC transparente 0,7 mm) · cinturão sensorial" if kind else f"lado {'direito' if side == 'D' else 'esquerdo'} da Espinha"
                    p = Pattern(pid, f"Painel {pid} · entre A{k - 1} (x {xa:.2f}) e A{k} (x {xb:.2f}) · {lab}", "cocoon", grp, L2, R2, outline)
                    base_u = ua_ if side == "D" else ub_; top_u = ub_ if side == "D" else ua_
                    p.edges = [("keder", path_uv(L2, R2, [(ua_, 0), (ub_, 0)])[:41]), ("keder", path_uv(L2, R2, [(ua_, 1), (ub_, 1)])[:41])]
                    if kind:
                        p.edges += [("ziper", [uv2d(L2, R2, base_u, v) for v in (0, 1)]), ("cristal", [uv2d(L2, R2, top_u, v) for v in (0, 1)])]
                        p.material = "PVC cristal 0,7 mm (ou ETFE 250 µm em 2 folhas) soldado por HF à lona; keder Ø10 nas laterais; zíper na base"
                        p.notes.append("Camadas por dentro (não fazem parte deste painel): tela mosquiteira fixa no trilho de harpão e cortina de voile + blackout em trilho curvo.")
                    else:
                        p.edges += [("clamp", [uv2d(L2, R2, top_u, v) for v in (0, 1)]), ("cristal" if clear else "keder_base", [uv2d(L2, R2, base_u, v) for v in (0, 1)])]
                    add_windows(p, xa, xb, L2, R2, ua_, ub_)
                    pats.append(p)
            continue
        if overlap:  # entalhe parcial da Espinha
            vs = (max(SX1, xa) - xa) / (xb - xa) if SX1 > xa else (min(SX2, xb) - xa) / (xb - xa)
            if SX1 > xa:   # entalhe aberto no lado R (A2)
                wps = [(0, 0), (1, 0), (1, 1), (ub, 1), (ub, vs), (ua, vs), (ua, 1), (0, 1)]
            else:          # entalhe aberto no lado L (A5)
                wps = [(0, 0), (ua, 0), (ua, vs), (ub, vs), (ub, 0), (1, 0), (1, 1), (0, 1)]
            outline = path_uv(L2, R2, wps)
            p = Pattern(f"P{k}", f"Painel P{k} · entre A{k - 1} (x {xa:.2f}) e A{k} (x {xb:.2f}) · com entalhe da Espinha de Luz", "cocoon", "externa", L2, R2, outline)
            p.edges = [("keder", L2), ("keder", R2), ("keder_base", [uv2d(L2, R2, 0, 0), uv2d(L2, R2, 0, 1)]), ("keder_base", [uv2d(L2, R2, 1, 0), uv2d(L2, R2, 1, 1)]),
                       ("clamp", path_uv(L2, R2, [(ua, vs), (ub, vs)])[:41] + [uv2d(L2, R2, ub, 1 if SX1 > xa else 0)]),
                       ("clamp", [uv2d(L2, R2, ua, vs), uv2d(L2, R2, ua, 1 if SX1 > xa else 0)])]
            add_windows(p, xa, xb, L2, R2)
            pats.append(p)
            continue
        p = Pattern(f"P{k}", f"Painel P{k} · {'cauda: de A7 (x 8,75) à ponta (x 9,56)' if tail else f'entre A{k - 1} (x {xa:.2f}) e A{k} (x {xb:.2f})'}", "cocoon", "externa", L2, R2)
        p.edges = [("keder", L2), ("bolsa_tubo" if tail else "keder", R2), ("keder_base", [uv2d(L2, R2, 0, 0), uv2d(L2, R2, 0, 1)]), ("keder_base", [uv2d(L2, R2, 1, 0), uv2d(L2, R2, 1, 1)])]
        if tail:
            vq = (9.3 - xa) / (xb - xa); p.lines.append(("clamp no quadro da cauda B08 (x 9,30)", path_uv(L2, R2, [(0, vq), (1, vq)])[:41]))
            p.notes.append("Ponta: a lona fecha em cone sobre o quadro da cauda; a borda R vira bolsa de cinta com olhal e o excedente é dobrado e clampado no anel Ø60,3.")
        add_windows(p, xa, xb, L2, R2)
        pats.append(p)
    # ---- Bico (P0): tira entre o anel A0 (trecho coberto) e a borda livre
    ths = C.bico_thetas(N)
    Lb = [C.section_point(C.X_FRONT, t) for t in ths]; Rb = [C.bico_point(t, 1.0) for t in ths]
    L2, R2 = unfold_strip(Lb, Rb)
    p = Pattern("P0", "Painel P0 · BICO · do anel A0 (θ ± 1,15 rad do topo) à borda livre em balanço", "cocoon", "externa", L2, R2)
    p.edges = [("keder", L2), ("bolsa_tubo", R2), ("livre", [L2[0], R2[0]]), ("livre", [L2[-1], R2[-1]])]
    p.lines.append(("cumeeira Ø114,3 · presilhas @300 (clamp contínuo)", path_uv(L2, R2, [(0.5, 0), (0.5, 1)])[:17]))
    p.lines.append(("costela Ø48,3 a 50 % (linha de presilhas)", path_uv(L2, R2, [(0, 0.5), (1, 0.5)])[:41]))
    p.notes.append("Bordas laterais (θ = π/2 ± 1,15) coincidem com o encontro da borda livre com o anel: reforço triangular 300 mm e olhal do tirante.")
    pats.append(p)
    for q in pats:
        if q.material: continue
        q.material = "PVDF tipo II 1050 g/m² (poliéster de alta tenacidade, PVC + laca PVDF), cor creme Zion, tratamento anti-fungo, M2 / classe B"
        q.notes.append("Sobras: keder 45 mm; bolsa de base 60 mm; recortes com anel de reforço soldado de 150 mm em PVC 900 g/m².")
    # ---- forro interno (F1..F8): superfície a -0,17 m, do vidro (x 0,90) ao quadro da cauda, acima do rodapé (z ≥ 0,15)
    bays_in = [(C.X_GLASS, C.ARCH_X[1])] + [(C.ARCH_X[i], C.ARCH_X[i + 1]) for i in range(1, 7)] + [(C.ARCH_X[7], 9.3)]

    def liner_curve(x, n=N):
        t0, t1 = C.theta_range(x); a, b = C.a(x) - 0.17, C.b(x) - 0.17
        pts = []
        for i in range(n + 1):
            t = t0 + (t1 - t0) * i / n; y = a * math.cos(t); z = C.ZC + b * math.sin(t)
            pts.append((C.shear(x, max(z, 0.15)), y, max(z, 0.15)))
        return pts
    for k, (xa, xb) in enumerate(bays_in, 1):
        L2, R2 = unfold_strip(liner_curve(xa), liner_curve(xb))
        xm = (xa + xb) / 2; ua = u_of_theta(xm, math.pi / 2 - HT); ub = u_of_theta(xm, math.pi / 2 + HT)
        overlap = max(xa, SX1) < min(xb, SX2)
        p = Pattern(f"F{k}", f"Forro F{k} · entre x {xa:.2f} e x {xb:.2f} (superfície interna, a 0,17 m da membrana)", "cocoon", "interna", L2, R2)
        p.edges = [("harpao", L2), ("harpao", R2), ("harpao", [uv2d(L2, R2, 0, 0), uv2d(L2, R2, 0, 1)]), ("harpao", [uv2d(L2, R2, 1, 0), uv2d(L2, R2, 1, 1)])]
        if overlap:
            v0 = (max(SX1, xa) - xa) / (xb - xa); v1 = (min(SX2, xb) - xa) / (xb - xa)
            p.holes.append(("abertura da Espinha de Luz (requadro)", path_uv(L2, R2, [(ua, v0), (ub, v0), (ub, v1), (ua, v1)]), "abertura"))
        add_windows(p, xa, xb, L2, R2)
        p.material = "Trevira CS (poliéster FR) 210 g/m² tensionado, cor creme; ou lona translúcida 350 g/m² sobre a Espinha"
        p.notes.append("Bordas com harpão soldado (perfil 40 x 20); recortes dos Olhos e dos difusores com anel de harpão próprio.")
        pats.append(p)
    return pats


# =============================================================================== SAFARI
def zenith_patterns():
    Z = Zenith(); pats = []; N = 96
    x0, x1, y0, y1 = Z.roof_bounds(); NS = 6
    ys = [y0 + (y1 - y0) * i / NS for i in range(NS + 1)]

    def line_y(y, n=N):
        return [(x0 + (x1 - x0) * i / n, y, Z.roof_z(x0 + (x1 - x0) * i / n, y)) for i in range(n + 1)]
    for k in range(NS):
        ya, yb = ys[k], ys[k + 1]
        L2, R2 = unfold_strip(line_y(ya), line_y(yb))
        p = Pattern(f"S{k + 1}", f"Faixa S{k + 1} · y {ya:.2f} a {yb:.2f} · de x {x0:.2f} (frente) a x {x1:.2f} (fundos)", "zenith", "externa", L2, R2)
        p.edges = [("cabo" if k == 0 else "solda", L2), ("cabo" if k == NS - 1 else "solda", R2), ("cabo", [uv2d(L2, R2, 0, 0), uv2d(L2, R2, 0, 1)]), ("cabo", [uv2d(L2, R2, 1, 0), uv2d(L2, R2, 1, 1)])]
        for pk in Z.PEAKS:
            if ya - pk["r"] < pk["y"] < yb + pk["r"]:
                pts = []
                for i in range(48):
                    a = 2 * math.pi * i / 48; px, py = pk["x"] + pk["r"] * math.cos(a), pk["y"] + pk["r"] * math.sin(a)
                    if ya <= py <= yb: pts.append(uv2d(L2, R2, (px - x0) / (x1 - x0), (py - ya) / (yb - ya)))
                if len(pts) > 3: p.holes.append((f'{pk["name"]} · anel Ø{2 * pk["r"]:.2f} (clamp)'.replace(".", ","), pts, "recorte"))
        # postes e cantos: marcas
        for (px, py) in Z.posts():
            if abs(py - ya) < 1e-6 or abs(py - yb) < 1e-6 or abs(px - x0) < 1e-6 or abs(px - x1) < 1e-6:
                if ya - 1e-6 <= py <= yb + 1e-6: p.lines.append((f"poste / chapa de canto ({px:.1f}; {py:.1f})", [uv2d(L2, R2, (px - x0) / (x1 - x0), (py - ya) / (yb - ya))]))
        p.material = "PVDF tipo II 1050 g/m², creme Zion; faixas soldadas em fábrica (HF 40 mm) numa peça única de ≈ 105 m² (fardo ≈ 120 kg)"
        p.notes.append("Form-finding obrigatório: a superfície de dois cumes é anticlástica entre os cumes; os padrões aqui são geométricos (forma-alvo) e devem ser ajustados pelo confeccionista com compensação biaxial.")
        pats.append(p)
    # forro: sobre o corpo, a -0,30 m, 4 faixas
    bx0, bx1, by0, by1 = 0.0, Z.L, -Z.W / 2, Z.W / 2; NL = 4
    ysl = [by0 + (by1 - by0) * i / NL for i in range(NL + 1)]

    def liner_y(y, n=N):
        return [(bx0 + (bx1 - bx0) * i / n, y, Z.liner_z(bx0 + (bx1 - bx0) * i / n, y)) for i in range(n + 1)]
    for k in range(NL):
        ya, yb = ysl[k], ysl[k + 1]
        L2, R2 = unfold_strip(liner_y(ya), liner_y(yb))
        p = Pattern(f"F{k + 1}", f"Forro F{k + 1} · y {ya:.2f} a {yb:.2f} · sobre o corpo (x 0 a {Z.L:.1f}), 0,30 m sob a membrana", "zenith", "interna", L2, R2)
        p.edges = [("harpao", L2), ("harpao", R2), ("harpao", [uv2d(L2, R2, 0, 0), uv2d(L2, R2, 0, 1)]), ("harpao", [uv2d(L2, R2, 1, 0), uv2d(L2, R2, 1, 1)])]
        for pk in Z.PEAKS:
            rr = pk["r"] + 0.25
            if ya - rr < pk["y"] < yb + rr:
                pts = []
                for i in range(48):
                    a = 2 * math.pi * i / 48; px, py = pk["x"] + rr * math.cos(a), pk["y"] + rr * math.sin(a)
                    if ya <= py <= yb and bx0 <= px <= bx1: pts.append(uv2d(L2, R2, (px - bx0) / (bx1 - bx0), (py - ya) / (yb - ya)))
                if len(pts) > 3: p.holes.append((f'poço de luz do {pk["name"].split(" (")[0].lower()} Ø{2 * rr:.2f}'.replace(".", ","), pts, "abertura"))
        p.material = "Trevira CS 210 g/m² tensionado, creme; harpão nos trilhos suspensos por cabos Ø4"
        pats.append(p)
    return pats


# =============================================================================== LODGE
def lodge_patterns(L=None):
    L = L or Lodge(); pats = []; N = 40
    ro = L.r_corner() + L.OVER / math.cos(math.pi / L.N)

    def r_edge(ang):
        k = (ang - math.pi / 8) % (2 * math.pi / L.N) - math.pi / L.N
        return ro * math.cos(math.pi / L.N) / math.cos(k)

    def radial(ang, r0, r1, zfn, n=N):
        return [(r0 + (r1 - r0) * i / n) * math.cos(ang) for i in range(n + 1)] and \
               [((r0 + (r1 - r0) * i / n) * math.cos(ang), (r0 + (r1 - r0) * i / n) * math.sin(ang), zfn(r0 + (r1 - r0) * i / n)) for i in range(n + 1)]
    for k in range(L.N):
        a0 = math.pi / 8 + 2 * math.pi * k / L.N; a1 = a0 + 2 * math.pi / L.N   # caibros nos vértices
        Lc = radial(a0, L.R_LANTERN, r_edge(a0), lambda r: L.roof_z(r) if r < r_edge(a0) - 1e-6 else L.Z_EAVE - 0.35)
        Rc = radial(a1, L.R_LANTERN, r_edge(a1), lambda r: L.roof_z(r) if r < r_edge(a1) - 1e-6 else L.Z_EAVE - 0.35)
        L2, R2 = unfold_strip(Lc, Rc)
        p = Pattern(f"G{k + 1}", f"Gomo G{k + 1} · do anel da lanterna (r {L.R_LANTERN:.2f}) ao beiral, entre os caibros R{k + 1} e R{(k + 1) % L.N + 1}", "lodge", "externa", L2, R2)
        p.edges = [("solda", L2), ("solda", R2), ("clamp", [uv2d(L2, R2, 0, 0), uv2d(L2, R2, 0, 1)]), ("cabo", path_uv(L2, R2, [(1, 0), (1, 1)])[:17])]
        p.lines.append(("caibro Ø76,1 sob a costura (presilhas @400)", L2))
        p.material = "PVDF tipo II 1050 g/m², creme Zion; 8 gomos soldados em fábrica numa peça única (≈ 73 m²) com abertura central Ø1,50"
        p.notes.append("Todos os gomos são iguais (octógono regular): confeccionar 1 gabarito e 8 peças; borda com bolsa de cabo Ø10 e pingadeira no perfil de beiral I01.")
        pats.append(p)
    # forro: 8 gomos do anel da lanterna (r 0,80) ao apótema (beiral interno), a -0,20 m
    for k in range(L.N):
        a0 = math.pi / 8 + 2 * math.pi * k / L.N; a1 = a0 + 2 * math.pi / L.N
        ap = L.F / 2 - 0.05
        fn = lambda r: L.roof_z(r) - 0.20
        Lc = radial(a0, L.R_LANTERN + 0.05, ap / math.cos(0), fn); Rc = radial(a1, L.R_LANTERN + 0.05, ap / math.cos(0), fn)
        L2, R2 = unfold_strip(Lc, Rc)
        p = Pattern(f"F{k + 1}", f"Forro F{k + 1} · gomo interno do anel da lanterna ao anel de beiral, 0,20 m sob a membrana", "lodge", "interna", L2, R2)
        p.edges = [("harpao", L2), ("harpao", R2), ("harpao", [uv2d(L2, R2, 0, 0), uv2d(L2, R2, 0, 1)]), ("harpao", [uv2d(L2, R2, 1, 0), uv2d(L2, R2, 1, 1)])]
        p.material = "Trevira CS 210 g/m² tensionado, creme; harpão nos trilhos dos caibros e no anel de beiral"
        pats.append(p)
    return pats


# =============================================================================== CÁPSULA (revestimento rígido)
def capsule_patterns():
    K = Capsule(); pats = []; N = 24
    SECT = [("T", 0.25 * math.pi, 0.75 * math.pi, "topo"), ("E", 0.75 * math.pi, 1.25 * math.pi, "lado esquerdo"), ("B", 1.25 * math.pi, 1.75 * math.pi, "barriga"), ("D", -0.25 * math.pi, 0.25 * math.pi, "lado direito")]
    xs = [0.75] + K.RINGS[1:] + [8.38]   # anéis + ponta da cauda

    def arc(x, t0, t1, inner=False, n=N):
        a, b = (K.A - K.SKIN, K.B - K.SKIN) if inner else (K.A, K.B); s = K.s(x)
        return [(x, *K._pt(t0 + (t1 - t0) * i / n, max(s, 0.02), a, b)) for i in range(n + 1)]
    for grp, inner, mat in (("revestimento", False, "ACM 4 mm (alumínio composto) curvado a frio, creme Zion, fixado nos anéis com perfis H de alumínio e EPDM"),
                            ("interna", True, "Compensado naval 12 mm curvado (2 x 6 mm colados), acabamento carvalho, fixado nas longarinas")):
        for k in range(len(xs) - 1):
            xa, xb = xs[k], xs[k + 1]
            for (sid, t0, t1, sname) in SECT:
                # painel glass? Anel de Luz (topo, x 3,95-4,40) e Visor (x < 1,20 acima do piso)
                if sid == "T" and xa < K.LIGHT_RING[1] and xb > K.LIGHT_RING[0] and not inner: continue
                if xb <= K.X_NOSE + 1e-6 and sid != "B": continue
                L2, R2 = unfold_strip(arc(xa, t0, t1, inner), arc(xb, t0, t1, inner))
                pid = f"{'R' if not inner else 'I'}{k + 1}{sid}"
                p = Pattern(pid, f"{'Revestimento externo' if not inner else 'Revestimento interno'} {pid} · {sname} · entre x {xa:.2f} e x {xb:.2f}", "capsule", grp, L2, R2, stations_every=0.25)
                p.edges = [("junta", L2), ("junta", R2), ("junta", [uv2d(L2, R2, 0, 0), uv2d(L2, R2, 0, 1)]), ("junta", [uv2d(L2, R2, 1, 0), uv2d(L2, R2, 1, 1)])]
                for ph in K.PORTHOLES:
                    if xa <= ph["x"] <= xb:
                        thc = math.pi if ph["side"] > 0 else 0.0
                        if not (t0 - 1e-6 <= thc <= t1 + 1e-6): continue
                        pts = []
                        for i in range(36):
                            a = 2 * math.pi * i / 36
                            px = ph["x"] + ph["r"] * math.cos(a); pz = ph["z"] + ph["r"] * math.sin(a)
                            # ângulo do ponto (y, z) na seção
                            th = thc + math.atan2(pz - K.ZC, 0.001) * 0 + (ph["r"] * math.sin(a)) / (K.B * 1.0) * (-1 if ph["side"] > 0 else 1)
                            pts.append(uv2d(L2, R2, (th - t0) / (t1 - t0), (px - xa) / (xb - xa)))
                        p.holes.append((f'{ph["name"]} Ø{2 * ph["r"]:.2f}'.replace(".", ","), pts, "recorte"))
                p.material = mat
                pats.append(p)
    return pats


PATTERNS = {"cocoon": cocoon_patterns, "zenith": zenith_patterns, "lodge": lodge_patterns, "capsule": capsule_patterns, "cocoon_s": lambda: cocoon_patterns(CocoonSensorial())}
TAG = {"cocoon": "ZC", "zenith": "ZS", "lodge": "ZL", "capsule": "ZK", "cocoon_s": "ZCS"}
NAME = {"cocoon": "ZION CASULO", "zenith": "ZION SAFARI", "lodge": "ZION LODGE 38", "capsule": "ZION CÁPSULA", "cocoon_s": "ZION CASULO SENSORIAL"}
GRP = {"externa": "LONA EXTERNA", "interna": "FORRO / REVESTIMENTO INTERNO", "revestimento": "REVESTIMENTO EXTERNO"}


def dedupe(pats):
    """agrupa padrões geometricamente iguais (mesmo grupo): mantém o primeiro com qty acumulada e lista dos códigos."""
    out = []; seen = {}
    for p in pats:
        key = (p.group, signature(p))
        if key in seen:
            q = seen[key]; q.qty += p.qty; q.notes = q.notes; q.same = getattr(q, "same", [q.id]) + [p.id]
        else:
            p.same = [p.id]; seen[key] = p; out.append(p)
    return out


# =============================================================================== pranchas
def fit(pat, box=(90, 130, 940, 780)):
    X0, Y0, X1, Y1 = box
    k = min((X1 - X0) / max(pat.w, 0.3), (Y1 - Y0) / max(pat.h, 0.3))
    k = min(k, 400)
    ox = X0 + ((X1 - X0) - pat.w * k) / 2; oy = Y1 - ((Y1 - Y0) - pat.h * k) / 2
    return k, ox, oy


def draw_pattern(sh, pat, box=(90, 130, 940, 780)):
    k, ox, oy = fit(pat, box)
    X = lambda x: ox + x * k; Y = lambda y: oy - y * k
    P = lambda pts: " ".join(f"{X(x):.1f},{Y(y):.1f}" for x, y in pts)
    # malha 0,50 m
    step = 0.5 if pat.w < 8 and pat.h < 8 else 1.0
    gx = 0.0
    while gx <= pat.w + 1e-6:
        sh.add(f'<line x1="{X(gx):.1f}" y1="{Y(0):.1f}" x2="{X(gx):.1f}" y2="{Y(pat.h):.1f}" stroke="{SAND}" stroke-width="0.5"/>')
        sh.text_px(X(gx), Y(0) + 12, f"{gx:.1f}".replace(".", ","), size=7.5, fill=EARTH); gx += step
    gy = 0.0
    while gy <= pat.h + 1e-6:
        sh.add(f'<line x1="{X(0):.1f}" y1="{Y(gy):.1f}" x2="{X(pat.w):.1f}" y2="{Y(gy):.1f}" stroke="{SAND}" stroke-width="0.5"/>')
        sh.text_px(X(0) - 6, Y(gy) + 3, f"{gy:.1f}".replace(".", ","), size=7.5, fill=EARTH, anchor="end"); gy += step
    sh.add(f'<polygon points="{P(pat.outline)}" fill="#F3EDE0" stroke="{GREEN}" stroke-width="1.2" stroke-linejoin="round"/>')
    for kind, pts in pat.edges:
        col, w, _ = KINDS[kind]
        sh.add(f'<polyline points="{P(pts)}" fill="none" stroke="{col}" stroke-width="{w}" stroke-linecap="round" stroke-linejoin="round"/>')
    for lab, pts in pat.lines:
        if len(pts) == 1:
            sh.add(f'<circle cx="{X(pts[0][0]):.1f}" cy="{Y(pts[0][1]):.1f}" r="4" fill="{STEEL}"/>'); sh.text_px(X(pts[0][0]) + 6, Y(pts[0][1]) + 3, lab, size=7.5, fill=STEEL, anchor="start")
        else:
            sh.add(f'<polyline points="{P(pts)}" fill="none" stroke="{STEEL}" stroke-width="1.0" stroke-dasharray="6 3"/>')
            mid = pts[len(pts) // 2]; sh.text_px(X(mid[0]) + 5, Y(mid[1]) - 4, lab, size=7.5, fill=STEEL, anchor="start")
    for lab, pts, kind in pat.holes:
        sh.add(f'<polygon points="{P(pts)}" fill="{CREAM}" stroke="{GREEN}" stroke-width="1.0" stroke-dasharray="4 2"/>')
        cx = sum(p[0] for p in pts) / len(pts); cy = sum(p[1] for p in pts) / len(pts)
        sh.text_px(X(cx), Y(cy) + 3, lab, size=7, fill=GREEN)
    # cotas gerais
    sh.add(f'<line x1="{X(0):.1f}" y1="{Y(0) + 26:.1f}" x2="{X(pat.w):.1f}" y2="{Y(0) + 26:.1f}" stroke="{EARTH}" stroke-width="0.8"/>')
    sh.text_px((X(0) + X(pat.w)) / 2, Y(0) + 38, f"largura máx. {pat.w:.3f} m".replace(".", ","), size=9, fill=EARTH)
    sh.add(f'<line x1="{X(pat.w) + 26:.1f}" y1="{Y(0):.1f}" x2="{X(pat.w) + 26:.1f}" y2="{Y(pat.h):.1f}" stroke="{EARTH}" stroke-width="0.8"/>')
    sh.text_px(X(pat.w) + 32, (Y(0) + Y(pat.h)) / 2, f"altura máx. {pat.h:.3f} m".replace(".", ","), size=9, fill=EARTH, anchor="start", rotate=90)
    # marcas L / R
    sh.text_px(X(pat.L2[0][0]) - 4, Y(pat.L2[0][1]) - 6, "L0", size=8, weight=700, anchor="end"); sh.text_px(X(pat.R2[0][0]) + 4, Y(pat.R2[0][1]) - 6, "R0", size=8, weight=700, anchor="start")
    sh.text_px(X(pat.L2[-1][0]) - 4, Y(pat.L2[-1][1]) - 6, "L1", size=8, weight=700, anchor="end"); sh.text_px(X(pat.R2[-1][0]) + 4, Y(pat.R2[-1][1]) - 6, "R1", size=8, weight=700, anchor="start")
    return k


def pattern_sheet(model, pat, code, folder):
    sh = Sheet(1600, 1000)
    sh.header(f"{NAME[model].title()} · Padrão de corte {pat.id}", f"{pat.name} · {GRP[pat.group]} · geometria de forma final (sem compensação) · metros")
    draw_pattern(sh, pat)
    # tabela de estações
    X0, Y0 = 1010, 118
    sh.text_px(X0, Y0, "TABELA DE ESTAÇÕES (ao longo da borda L, a cada " + f"{pat.stations_every:.2f}".replace(".", ",") + " m)", size=9.5, weight=700, spacing=0.18, anchor="start")
    cols = [0, 50, 110, 170, 230, 290, 360]
    for hd, cx in zip(["s (m)", "xL", "yL", "xR", "yR", "largura"], cols): sh.text_px(X0 + cx, Y0 + 20, hd, size=7.5, fill=EARTH, spacing=0.12, anchor="start")
    sh.add(f'<line x1="{X0}" y1="{Y0 + 26}" x2="{X0 + 420}" y2="{Y0 + 26}" stroke="{GREEN}" stroke-width="0.6"/>')
    rows = pat.stations(); yy = Y0 + 40; maxrows = 30
    if len(rows) > maxrows:   # amostra uniforme se a tira for longa
        idx = [round(i * (len(rows) - 1) / (maxrows - 1)) for i in range(maxrows)]; rows = [rows[i] for i in idx]
    for r in rows:
        for v, cx in zip(r, cols): sh.text_px(X0 + cx, yy, f"{v:.3f}".replace(".", ","), size=7.8, anchor="start")
        yy += 12.5
    yy += 8
    sh.text_px(X0, yy, "DADOS DO PAINEL", size=9.5, weight=700, spacing=0.18, anchor="start"); yy += 16
    for lab, val in [("Quantidade", f"{pat.qty} (códigos: {', '.join(getattr(pat, 'same', [pat.id]))})"), ("Área de forma final", f"{pat.area:.2f} m²".replace(".", ",")), ("Comprimento L / R", f"{pat.lenL:.2f} / {pat.lenR:.2f} m".replace(".", ",")),
                     ("Caixa (larg. x alt.)", f"{pat.w:.2f} x {pat.h:.2f} m".replace(".", ",")), ("Material", pat.material)]:
        sh.text_px(X0, yy, lab.upper(), size=7, fill=EARTH, spacing=0.14, anchor="start"); yy += 10
        for line in wrap(val, 78): sh.text_px(X0, yy, line, size=8.2, anchor="start"); yy += 11
        yy += 3
    # legenda das bordas presentes
    yy += 6; sh.text_px(X0, yy, "BORDAS E RECORTES", size=9.5, weight=700, spacing=0.18, anchor="start"); yy += 14
    kinds = []
    for kind, _ in pat.edges:
        if kind not in kinds: kinds.append(kind)
    for kind in kinds:
        col, w, lab = KINDS[kind]
        sh.add(f'<line x1="{X0}" y1="{yy - 3}" x2="{X0 + 26}" y2="{yy - 3}" stroke="{col}" stroke-width="{w}"/>')
        for line in wrap(lab, 70): sh.text_px(X0 + 34, yy, line, size=7.8, anchor="start"); yy += 10.5
        yy += 2
    if pat.holes:
        sh.add(f'<rect x="{X0}" y="{yy - 10}" width="26" height="8" fill="{CREAM}" stroke="{GREEN}" stroke-width="0.8" stroke-dasharray="3 2"/>')
        sh.text_px(X0 + 34, yy, "RECORTE com anel de reforço soldado (150 mm) e clamp / requadro", size=7.8, anchor="start"); yy += 12
    # notas
    yy = 830
    sh.text_px(60, yy, "NOTAS", size=9.5, weight=700, spacing=0.18, anchor="start"); yy += 14
    notes = ["Padrão = geometria de forma final da superfície entre as costuras; o confeccionista aplica a compensação do tecido (ensaio biaxial; típico 0,5 a 1,5 % no urdume e 2 a 4 % na trama) e as sobras de bolsa/keder.",
             "Origem (0; 0) no canto inferior esquerdo da caixa; eixo x horizontal, y vertical; L0 → L1 é a borda esquerda (costura de referência), R0 → R1 a direita; estações medidas ao longo de L.",
             "Marcar em cada painel: código, modelo, orientação (L/R, base), linhas de presilha e centros dos recortes; etiqueta de tecido cosida no verso."] + pat.notes
    for n in notes:
        for line in wrap("· " + n, 150): sh.text_px(60, yy, line, size=8, anchor="start"); yy += 10.5
        yy += 1
    sh.title_block(NAME[model], f"Padrão de corte {pat.id}", "ver malha (m)", code, GRP[pat.group].title())
    sh.save(os.path.join(folder, f"{code}_{pat.id}.svg"))


def wrap(t, n):
    w = str(t).split(); out = []; cur = ""
    for x in w:
        if len(cur) + len(x) + 1 > n: out.append(cur); cur = x
        else: cur = (cur + " " + x).strip()
    return out + [cur]


def map_sheet(model, pats, folder):
    """mapa dos painéis: projeção simples do modelo com as costuras e os códigos."""
    sh = Sheet(1600, 1000)
    ext = [p for p in pats if p.group != "interna"]; inn = [p for p in pats if p.group == "interna"]
    sh.header(f"{NAME[model].title()} · Mapa dos painéis de lona e revestimento", f"{len(ext)} padrões externos · {len(inn)} padrões internos · costuras sobre a estrutura · códigos usados nas pranchas LN e no DXF")
    if model.startswith("cocoon"):
        C = CocoonSensorial() if model == "cocoon_s" else Cocoon(); sh.s, sh.ox, sh.oy = 72, 420, 560
        from drawings_cocoon import top_profile, bottom_profile, bico_side
        top = top_profile(); bot = bottom_profile(); sh.poly(top + bot[::-1], fill="#F3EDE0", stroke=GREEN, sw=1.4); sh.poly(bico_side(), fill="#F3EDE0", stroke=GREEN, sw=1.4)
        for i, x in enumerate(C.ARCH_X): sh.line(C.shear(x, 0), 0, C.shear(x, C.top(x)), C.top(x), "#1B2117", 2.2); sh.text(C.shear(x, C.top(x)), C.top(x) + 0.25, f"A{i}", 10, GREEN, weight=700)
        sh.line(C.shear(9.3, 0), 0, C.shear(9.3, C.top(9.3)), C.top(9.3), STEEL, 1.4, dash="4 3")
        for i in range(8):
            xa, xb = C.ARCH_X[i], (C.ARCH_X[i + 1] if i < 7 else 9.56); xm = (xa + xb) / 2
            lab = f"P{i + 1}" + ("D/E" if 2 <= i + 1 <= 5 and i + 1 in (3, 4, 5) else "")
            sh.text(C.shear(xm, 1.0), 1.0, lab, 12, GREEN, weight=700); sh.text(C.shear(xm, 0.6), 0.6, f"F{i + 1}", 10, "#6B7C84")
        sh.text(-1.3, 3.3, "P0 · BICO", 12, GREEN, weight=700)
        x1, x2, _ = C.SPINE; sh.poly([(C.shear(x, C.top(x)), C.top(x)) for x in np.linspace(x1, x2, 20)], close=False, stroke="#8B714E", sw=5, opacity=0.7); sh.text(4.2, 4.55, "Espinha de Luz: clamp nas treliças (P3, P4, P5 divididos em D/E; P2 e P6 com entalhe)", 9, EARTH)
        for w in C.WINDOWS:
            if w["tc"] > math.pi / 2: continue
            sh.poly([(x, z) for (x, y, z) in C.window_outline(w, 40)], fill=CREAM, stroke=GREEN, sw=0.8, dash="3 2")
        if hasattr(C, "CLEAR"):
            cl = C.CLEAR; sh.poly([(C.shear(cl["x1"], 0), 0), (C.shear(cl["x2"], 0), 0), (C.shear(cl["x2"], cl["z"]), cl["z"]), (C.shear(cl["x1"], cl["z"]), cl["z"])], fill="#BFD8E8", stroke="#4E8BB8", sw=1.4, opacity=0.9)
            sh.text((cl["x1"] + cl["x2"]) / 2, 1.3, "T3 / T4 · CRISTAL (zíper na base)", 9, "#2F5F80", weight=700)
        sh.line(-4, 0, 11, 0, GREEN, 1.0); sh.text(3.5, -0.45, "trilho de base 100 x 50 com perfil de arremate E02 (keder Ø13 + calha)", 9, "#4E6E8B")
        sh.text(5.0, -1.0, "Vista lateral direita · costuras sobre os arcos A0 a A7 (perfil duplo keder) · painéis externos P0 a P8 · forros F1 a F8 (harpão)", 10, GREEN)
    elif model == "zenith":
        Z = Zenith(); sh.s, sh.ox, sh.oy = 62, 380, 600; x0, x1, y0, y1 = Z.roof_bounds()
        sh.rect(x0, y0, x1, y1, fill="#F3EDE0", stroke=GREEN, sw=1.4); sh.rect(0, -Z.W / 2, Z.L, Z.W / 2, fill="none", stroke=GREEN, sw=0.8, dash="6 3")
        for i in range(1, 6): y = y0 + (y1 - y0) * i / 6; sh.line(x0, y, x1, y, "#C2734A", 1.4); sh.text(x1 + 0.4, y, f"solda S{i}/S{i + 1}", 8, "#C2734A", anchor="start")
        for i in range(6): sh.text((x0 + x1) / 2, y0 + (y1 - y0) * (i + 0.5) / 6, f"S{i + 1}", 12, GREEN, weight=700)
        for i in range(1, 4): y = -Z.W / 2 + Z.W * i / 4; sh.line(0, y, Z.L, y, "#6B7C84", 1.0, dash="5 3")
        for i in range(4): sh.text(Z.L / 2, -Z.W / 2 + Z.W * (i + 0.5) / 4 - 0.3, f"F{i + 1}", 9, "#6B7C84")
        for p in Z.PEAKS: sh.circle(p["x"], p["y"], p["r"], fill=CREAM, stroke=GREEN, sw=1.0); sh.text(p["x"], p["y"] + p["r"] + 0.3, p["name"].split(" (")[0] + " · clamp", 8, GREEN)
        for (px, py) in Z.posts(): sh.circle(px, py, 0.12, fill=STEEL, stroke="none")
        sh.text((x0 + x1) / 2, y0 - 0.7, "Planta da cobertura · 6 faixas longitudinais soldadas (HF) · perímetro em bolsa de cabo Ø12 com chapas de canto nos 7 postes · forro em 4 faixas sobre o corpo", 10, GREEN)
    elif model == "lodge":
        L = Lodge(); sh.s, sh.ox, sh.oy = 80, 800, 520
        ro = L.r_corner() + L.OVER / math.cos(math.pi / L.N)
        V = L.vertices(ro); sh.poly(V, fill="#F3EDE0", stroke=GREEN, sw=1.4); sh.poly(L.vertices(), fill="none", stroke=GREEN, sw=0.8, dash="6 3")
        sh.circle(0, 0, L.R_LANTERN, fill=CREAM, stroke=GREEN, sw=1.2); sh.text(0, L.R_LANTERN + 0.3, "anel da lanterna Ø1,50 · clamp I02", 8, GREEN)
        for k in range(L.N):
            a = math.pi / 8 + 2 * math.pi * k / L.N; r = ro * math.cos(math.pi / L.N) / math.cos(0)
            sh.line(L.R_LANTERN * math.cos(a), L.R_LANTERN * math.sin(a), ro * math.cos(a), ro * math.sin(a), "#C2734A", 1.4)
            am = a + math.pi / L.N; rm = (L.R_LANTERN + ro) / 2
            sh.text(rm * math.cos(am), rm * math.sin(am), f"G{k + 1}", 12, GREEN, weight=700); sh.text(rm * 0.65 * math.cos(am), rm * 0.65 * math.sin(am), f"F{k + 1}", 9, "#6B7C84")
        sh.text(0, -ro - 0.6, "Planta da cobertura · 8 gomos iguais soldados sobre os caibros · bolsa de cabo no beiral (perfil I01) · forro em 8 gomos (harpão)", 10, GREEN)
    else:
        K = Capsule(); sh.s, sh.ox, sh.oy = 110, 260, 560
        sh.poly(K.profile(), fill="#F3EDE0", stroke=GREEN, sw=1.4)
        for i, x in enumerate(K.RINGS): sh.line(x, K.bottom(x), x, K.top(x), "#1B2117", 2.0); sh.text(x, K.top(x) + 0.2, f"R{i + 1}", 8, GREEN, weight=700)
        for k in range(len(K.RINGS)):
            xa = K.RINGS[k]; xb = K.RINGS[k + 1] if k + 1 < len(K.RINGS) else 8.38
            sh.text((xa + xb) / 2, K.ZC + 1.0, f"{k + 1}T", 8, GREEN); sh.text((xa + xb) / 2, K.ZC, f"{k + 1}D/E", 8, GREEN); sh.text((xa + xb) / 2, K.ZC - 1.0, f"{k + 1}B", 8, GREEN)
        x1, x2 = K.LIGHT_RING; sh.rect(x1, K.top(x1) - 0.1, x2, K.top(x1) + 0.05, fill=GLASS, stroke=GREEN, sw=0.8); sh.text((x1 + x2) / 2, K.top(x1) + 0.45, "Anel de Luz (vidro)", 8, GREEN)
        sh.poly([(x, K.top(x)) for x in np.linspace(0.01, K.X_NOSE, 20)] + [(K.X_NOSE, K.bottom(K.X_NOSE))] + [(x, K.bottom(x)) for x in np.linspace(K.X_NOSE, 0.01, 20)], fill=GLASS, stroke=GREEN, sw=0.8, opacity=0.6); sh.text(0.6, K.ZC, "Visor", 8, GREEN)
        sh.text(4.2, K.bottom(4.2) - 0.6, "Vista lateral · painéis de ACM entre anéis em 4 setores (T topo, D/E lados, B barriga) · juntas em perfil H · revestimento interno em compensado curvado", 10, GREEN)
    # legenda de bordas
    X0, Y0 = 1080, 120
    sh.text_px(X0, Y0, "TIPOS DE BORDA (ver LN-90)", size=10, weight=700, spacing=0.22, anchor="start"); yy = Y0 + 18
    used = []
    for p in pats:
        for kind, _ in p.edges:
            if kind not in used: used.append(kind)
    for kind in used:
        col, w, lab = KINDS[kind]
        sh.add(f'<line x1="{X0}" y1="{yy - 3}" x2="{X0 + 30}" y2="{yy - 3}" stroke="{col}" stroke-width="{w}"/>')
        for line in wrap(lab, 62): sh.text_px(X0 + 40, yy, line, size=8.5, anchor="start"); yy += 11
        yy += 3
    yy += 10; sh.text_px(X0, yy, "QUADRO DE PAINÉIS", size=10, weight=700, spacing=0.22, anchor="start"); yy += 16
    for hd, cx in zip(["Cód.", "Qtd", "Caixa (m)", "Área m²", "Grupo"], [0, 60, 100, 220, 290]): sh.text_px(X0 + cx, yy, hd, size=7.5, fill=EARTH, spacing=0.12, anchor="start")
    yy += 5; sh.add(f'<line x1="{X0}" y1="{yy}" x2="{X0 + 440}" y2="{yy}" stroke="{GREEN}" stroke-width="0.6"/>'); yy += 12
    tot = {"externa": 0, "interna": 0, "revestimento": 0}
    for p in pats:
        tot[p.group] += p.area * p.qty
        if yy > 960: continue
        for v, cx in zip([p.id, str(p.qty), f"{p.w:.2f} x {p.h:.2f}".replace(".", ","), f"{p.area:.2f}".replace(".", ","), GRP[p.group].split(" /")[0].title()], [0, 60, 100, 220, 290]): sh.text_px(X0 + cx, yy, v, size=7.8, anchor="start")
        yy += 11
    yy += 8
    for g, a in tot.items():
        if a: sh.text_px(X0, yy, f"Total {GRP[g].lower()}: {a:.1f} m² de forma final (+ 12 a 15 % de sobras, bolsas e emendas)".replace(".", ","), size=8.5, weight=600, anchor="start"); yy += 12
    sh.title_block(NAME[model], "Mapa dos painéis de lona", "s/ escala", "LN-01", "Costuras sobre a estrutura · códigos dos padrões")
    sh.save(os.path.join(folder, "LN-01_mapa_paineis.svg"))


def fix_sheet(model, folder):
    """LN-90 · como a lona se prende aos ferros: seis detalhes esquemáticos (1:5) com as notas do modelo."""
    sh = Sheet(1600, 1000)
    sh.header(f"{NAME[model].title()} · Fixação da lona nos ferros", "Detalhes esquemáticos 1:5 · keder, bolsa de base, bolsa de tubo, clamp de anel, harpão do forro, bolsa de cabo · sequência de tensionamento")
    k = 1.5   # px por mm
    def panel(X0, Y0, title, draw):
        sh.add(f'<rect x="{X0}" y="{Y0}" width="470" height="330" fill="none" stroke="{GREEN}" stroke-width="0.8"/>')
        sh.text_px(X0 + 10, Y0 + 18, title, size=10, weight=700, spacing=0.14, anchor="start")
        draw(X0 + 200, Y0 + 190)
    def tube(cx, cy, d_mm, col=STEEL):
        sh.add(f'<circle cx="{cx:.1f}" cy="{cy:.1f}" r="{d_mm / 2 * k:.1f}" fill="{col}" stroke="{GREEN}" stroke-width="0.8"/>')
    def box(x, y, w, h, col="#BDBDB8"):
        sh.add(f'<rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" fill="{col}" stroke="{GREEN}" stroke-width="0.8"/>')
    def memb(pts, col="#8B714E"):
        sh.add(f'<polyline points="{" ".join(f"{x:.1f},{y:.1f}" for x, y in pts)}" fill="none" stroke="{col}" stroke-width="3" stroke-linecap="round" stroke-linejoin="round"/>')
    def lab(x, y, t, a="start", col=GREEN):
        sh.text_px(x, y, t, size=8, anchor=a, fill=col)
    def note(X0, Y0, t):
        for i, line in enumerate(wrap(t, 88)): sh.text_px(X0 + 10, Y0 + 300 + i * 11, line, size=7.8, anchor="start", fill=EARTH)
    # 1 keder duplo no arco
    def d1(cx, cy):
        tube(cx, cy + 30, 88.9)
        box(cx - 35 * k, cy - 36 - 24 * k, 70 * k, 24 * k)
        for s_ in (-1, 1): sh.add(f'<circle cx="{cx + s_ * 22 * k:.1f}" cy="{cy - 36 - 12 * k:.1f}" r="{5 * k}" fill="#F3EDE0" stroke="{GREEN}" stroke-width="0.8"/>')
        memb([(cx - 22 * k - 6, cy - 36 - 12 * k), (cx - 150, cy - 20)]); memb([(cx + 22 * k + 6, cy - 36 - 12 * k), (cx + 150, cy - 20)])
        sh.add(f'<line x1="{cx}" y1="{cy - 36}" x2="{cx}" y2="{cy - 36 + 30}" stroke="{GREEN}" stroke-width="2.5"/>')
        lab(cx + 75, cy + 35, "arco Ø88,9 (Ø101,6 no A0)"); lab(cx - 190, cy - 110, "perfil duplo keder alu 70 x 24 (E01)"); lab(cx - 190, cy - 96, "cordão keder Ø10 soldado na lona"); lab(cx + 75, cy - 12, "presilha inox 3 mm + M8 @300 (E03)")
        lab(cx - 190, cy + 60, "membrana P(n)", col="#8B714E"); lab(cx + 110, cy + 60, "membrana P(n+1)", col="#8B714E")
    # 2 bolsa de base
    def d2(cx, cy):
        box(cx - 50 * k, cy + 10, 100 * k, 50 * k)
        box(cx - 50 * k, cy + 10 - 60 * k, 40 * k, 60 * k, "#D9D4C7")
        sh.add(f'<circle cx="{cx - 30 * k:.1f}" cy="{cy + 10 - 48 * k:.1f}" r="{6.5 * k}" fill="#F3EDE0" stroke="{GREEN}" stroke-width="0.8"/>')
        memb([(cx - 30 * k - 8, cy + 10 - 48 * k), (cx - 30 * k - 40, cy + 10 - 60 * k - 70)])
        sh.add(f'<line x1="{cx - 50 * k:.1f}" y1="{cy + 10 - 60 * k:.1f}" x2="{cx - 50 * k:.1f}" y2="{cy + 10 - 60 * k - 26:.1f}" stroke="{GREEN}" stroke-width="1.5"/>')
        lab(cx + 10, cy + 60, "trilho de base 100 x 50 x 3 (A07/A08)"); lab(cx + 10, cy + 74, "sobre a viga de borda U 150"); lab(cx - 20, cy - 60, "perfil de arremate E02:"); lab(cx - 20, cy - 46, "keder Ø13 + calha 80 mm"); lab(cx - 190, cy - 110, "membrana desce por dentro", col="#8B714E")
    # 3 bolsa de tubo
    def d3(cx, cy):
        r = 30.15 * k; tube(cx, cy - 20, 60.3)
        sh.add(f'<path d="M{cx - r - 6:.1f},{cy - 20:.1f} A{r + 6:.1f},{r + 6:.1f} 0 1 1 {cx + r + 6:.1f},{cy - 20:.1f}" fill="none" stroke="#8B714E" stroke-width="3"/>')
        memb([(cx - r - 6, cy - 20), (cx - r - 6, cy + 40), (cx - 150, cy + 90)]); memb([(cx + r + 6, cy - 20), (cx + r + 6, cy + 40), (cx - r - 6, cy + 40)])
        lab(cx + 60, cy - 30, "tubo de borda Ø60,3"); lab(cx + 60, cy - 16, "(Bico B10 / quadro da cauda B08)"); lab(cx + 60, cy + 50, "bolsa fechada com solda HF"); lab(cx + 60, cy + 64, "+ fita keder no fecho"); lab(cx - 190, cy + 96, "aba de 60 mm cosida à face interna", col="#8B714E")
    # 4 clamp de anel
    def d4(cx, cy):
        box(cx - 110, cy, 220, 10 * k); box(cx - 110, cy - 6 * k - 6, 220, 6 * k, "#D9D4C7")
        memb([(cx - 180, cy - 34), (cx - 110, cy - 4), (cx + 110, cy - 4)])
        for x in (cx - 70, cx, cx + 70): sh.add(f'<rect x="{x - 3}" y="{cy - 6 * k - 34}" width="6" height="{6 * k + 34 + 10 * k}" fill="{STEEL}"/>')
        lab(cx + 120, cy + 12, "anel / chapa (cume, lanterna, Olho E06)"); lab(cx + 120, cy - 12, "barra alu 40 x 6 + EPDM 3 mm"); lab(cx - 190, cy - 56, "lona com reforço 150 mm", col="#8B714E"); lab(cx - 190, cy - 90, "parafusos M8 inox @150")
    # 5 harpão
    def d5(cx, cy):
        box(cx - 20 * k, cy - 40 * k, 40 * k, 20 * k)
        memb([(cx - 20 * k + 5, cy - 22 * k), (cx - 20 * k + 5, cy - 34 * k), (cx - 20 * k + 24, cy - 30 * k)], "#6B7C84"); memb([(cx - 20 * k + 5, cy - 22 * k), (cx - 170, cy + 50)], "#6B7C84")
        tube(cx, cy - 40 * k - 40, 48.3)
        lab(cx + 45, cy - 30 * k, "trilho de harpão alu 40 x 20 (E04)"); lab(cx + 45, cy - 40 * k - 40, "terça Ø48,3 / cabo"); lab(cx - 190, cy + 76, "harpão de PVC soldado na borda do forro", col="#6B7C84")
    # 6 bolsa de cabo
    def d6(cx, cy):
        r = 6 * k; sh.add(f'<circle cx="{cx}" cy="{cy}" r="{r}" fill="{STEEL}" stroke="{GREEN}" stroke-width="0.8"/>')
        sh.add(f'<path d="M{cx - r - 6:.1f},{cy:.1f} A{r + 6:.1f},{r + 6:.1f} 0 1 1 {cx + r + 6:.1f},{cy:.1f}" fill="none" stroke="#8B714E" stroke-width="3"/>')
        memb([(cx - r - 6, cy), (cx - r - 6, cy + 40), (cx - 170, cy + 60)]); memb([(cx + r + 6, cy), (cx + r + 6, cy + 40), (cx - r - 6, cy + 40)])
        lab(cx + 30, cy - 6, "cabo inox Ø12 (Safari) / Ø10 (Lodge)"); lab(cx + 30, cy + 8, "em bolsa contínua"); lab(cx + 30, cy + 56, "esticador M20 + chapa de canto D14"); lab(cx - 190, cy + 90, "flecha da catenária ± 20 mm", col="#8B714E")
    NOTES = ["Painéis deslizam no perfil da cauda para a frente; a tensão transversal vem das bolsas de base; presilhas a cada 300 mm.",
             "Esticadores M12 nos olhais do trilho, alternados (em cruz), até a pré-tensão de projeto (≈ 2,5 kN/m); calha drena para os tubos de queda.",
             "A bolsa é enfiada no tubo antes de fixar o tubo à cumeeira / ao quadro; sem furos na lona.",
             "Furar a lona só através do reforço; selar com fita butílica sob a barra; torque 6 N·m.",
             "Tensionar com espátula e soprador térmico (60 °C); sem rugas; recortes com anel de harpão.",
             "Cantos: chapa inox 8 mm com 3 furos (cabo, cabo, cinta) no olhal do poste; tensionar em cruz."]
    ps = [("1 · KEDER DUPLO NO ARCO / CAIBRO", d1), ("2 · BOLSA DE BASE NO TRILHO", d2), ("3 · BOLSA DE TUBO (BICO E CAUDA)", d3), ("4 · CLAMP DE ANEL (CUME, LANTERNA, OLHOS, ESPINHA)", d4), ("5 · HARPÃO DO FORRO INTERNO", d5), ("6 · BOLSA DE CABO DE BORDA", d6)]
    for i, (t, fn) in enumerate(ps):
        X0, Y0 = 60 + (i % 3) * 500, 110 + (i // 3) * 350
        panel(X0, Y0, t, fn); note(X0, Y0, NOTES[i])
    notes = {"cocoon_s": "Casulo Sensorial: como o Casulo, mais o cinturão transparente T3/T4: PVC cristal soldado por HF à lona na linha da terça (z 1,85), keder nos arcos, zíper #10 na base para abrir a parede; por dentro, tela mosquiteira no harpão e cortina de voile + blackout em trilho curvo.",
             "cocoon": "Casulo: painéis P1 a P8 no keder dos arcos (1) e nas bolsas de base (2); Bico P0 no keder do A0 e na bolsa do tubo de borda (3), presilhas na cumeeira e na costela; Espinha e Olhos em clamp (4); forro F1 a F8 em harpão (5).",
             "zenith": "Safari: peça única içada pelos anéis dos cumes e presa nos clamps (4); perímetro em bolsa de cabo Ø12 (6) com chapas de canto nos 7 postes; forro F1 a F4 em harpão (5) nos trilhos suspensos.",
             "lodge": "Lodge: peça única de 8 gomos içada pelo anel da lanterna (clamp 4, perfil I02), passada sobre o anel de beiral (perfil de borda I01) e tensionada pela bolsa de cabo Ø10 (6); presilhas nos caibros; forro em harpão (5).",
             "capsule": "Cápsula: revestimento rígido (ACM) em perfis H de alumínio com EPDM nos anéis; sem lona tensionada. Este quadro vale para a vela e para o forro têxtil opcional."}[model]
    yy = 830
    for line in wrap(notes, 190): sh.text_px(60, yy, line, size=9, anchor="start"); yy += 12
    seq = "SEQUÊNCIA: 1 conferir geometria da estrutura (topo dos arcos ± 15 mm) · 2 desenrolar e posicionar sem arrastar (proteger arestas) · 3 enfiar keder / cabos / bolsas sem tensão · 4 fechar clamps e anéis · 5 tensionar em cruz, em 3 passes (30 %, 70 %, 100 %) · 6 medir pré-tensão e flechas, registrar · 7 teste de água · vento > 30 km/h suspende."
    for line in wrap(seq, 190): sh.text_px(60, yy, line, size=9, anchor="start", fill=EARTH); yy += 12
    sh.title_block(NAME[model], "Fixação da lona nos ferros", "1:5 (esquemático)", "LN-90", "Keder · bolsas · clamps · harpão · cabos")
    sh.save(os.path.join(folder, "LN-90_fixacoes.svg"))


# =============================================================================== DXF e XLSX
def export_dxf(model, pats, path):
    import ezdxf
    doc = ezdxf.new("R2010"); doc.header["$INSUNITS"] = 6
    for name, col in (("CONTORNO", 7), ("KEDER", 5), ("BOLSA", 30), ("CLAMP", 1), ("RECORTE", 3), ("LINHA", 8), ("TEXTO", 2), ("MALHA", 253)):
        doc.layers.add(name, color=col)
    msp = doc.modelspace(); ox = 0.0
    for p in pats:
        msp.add_lwpolyline([(x + ox, y) for x, y in p.outline], close=True, dxfattribs={"layer": "CONTORNO"})
        for kind, pts in p.edges:
            lay = {"keder": "KEDER", "keder_base": "KEDER", "harpao": "KEDER", "bolsa_tubo": "BOLSA", "cabo": "BOLSA", "clamp": "CLAMP", "junta": "CLAMP", "solda": "LINHA", "livre": "LINHA", "cristal": "LINHA", "ziper": "KEDER"}[kind]
            msp.add_lwpolyline([(x + ox, y) for x, y in pts], dxfattribs={"layer": lay})
        for lab, pts, kind in p.holes:
            msp.add_lwpolyline([(x + ox, y) for x, y in pts], close=True, dxfattribs={"layer": "RECORTE"})
            cx = sum(q[0] for q in pts) / len(pts) + ox; cy = sum(q[1] for q in pts) / len(pts)
            msp.add_text(lab, dxfattribs={"layer": "TEXTO", "height": 0.05}).set_placement((cx, cy))
        for lab, pts in p.lines:
            if len(pts) > 1: msp.add_lwpolyline([(x + ox, y) for x, y in pts], dxfattribs={"layer": "LINHA"})
        msp.add_text(f"{p.id} · x{p.qty} · {p.area:.2f} m2", dxfattribs={"layer": "TEXTO", "height": 0.12}).set_placement((ox, -0.35))
        msp.add_text(f"{NAME[model]} · {GRP[p.group]} · {p.material[:60]}", dxfattribs={"layer": "TEXTO", "height": 0.06}).set_placement((ox, -0.55))
        ox += p.w + 0.6
    doc.saveas(path)


def export_xlsx(model, pats, path):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Quadro"
    ws.append(["Código", "Qtd", "Grupo", "Nome", "Largura (m)", "Altura (m)", "Área (m²)", "Comp. L (m)", "Comp. R (m)", "Material"])
    for p in pats: ws.append([p.id, p.qty, GRP[p.group], p.name, round(p.w, 3), round(p.h, 3), round(p.area, 3), round(p.lenL, 3), round(p.lenR, 3), p.material])
    for c in ws[1]: c.font = Font(bold=True)
    for p in pats:
        w = wb.create_sheet(p.id[:30]); w.append([p.name]); w.append(["s (m)", "xL", "yL", "xR", "yR", "largura (m)"])
        for r in p.stations(): w.append([round(v, 4) for v in r])
        w.append([]); w.append(["Contorno (x, y) em metros"])
        for x, y in p.outline: w.append([round(x, 4), round(y, 4)])
        for lab, pts, kind in p.holes:
            w.append([]); w.append([f"Recorte: {lab}"])
            for x, y in pts: w.append([round(x, 4), round(y, 4)])
    wb.save(path)


def build(models):
    for model in models:
        folder = os.path.join(ROOT, model, "lona"); os.makedirs(folder, exist_ok=True)
        for f in os.listdir(folder):
            if f.startswith("LN-") and f.endswith(".svg"): os.remove(os.path.join(folder, f))
        pats = dedupe(PATTERNS[model]())
        map_sheet(model, pats, folder)
        for i, p in enumerate(pats, 1):
            pattern_sheet(model, p, f"LN-{10 + i:02d}", folder)
        fix_sheet(model, folder)
        export_dxf(model, pats, os.path.join(folder, f"{TAG[model]}-LON-001_padroes.dxf"))
        export_xlsx(model, pats, os.path.join(folder, f"{TAG[model]}-LON-001_coordenadas.xlsx"))
        ext = sum(p.area * p.qty for p in pats if p.group != "interna"); inn = sum(p.area * p.qty for p in pats if p.group == "interna")
        print(f"{model}: {len(pats)} padrões · externa {ext:.1f} m² · interna {inn:.1f} m² → {folder}")


def sheets(model):
    folder = os.path.join(ROOT, model, "lona")
    return sorted(f for f in os.listdir(folder) if f.startswith("LN-") and f.endswith(".svg")) if os.path.isdir(folder) else []


if __name__ == "__main__":
    build(sys.argv[1:] or ["cocoon", "zenith", "lodge", "capsule"])
